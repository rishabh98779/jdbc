#!/usr/bin/env python3
"""Build the Java Concepts Integrated Visual Notes book.

    python3 build.py            # HTML + PDF + DOCX
    python3 build.py html       # HTML only (fast iteration)
    python3 build.py pdf
    python3 build.py docx
"""
from __future__ import annotations

import os
import re
import sys
from html.parser import HTMLParser

ROOT = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(ROOT, "src")
OUT = os.path.join(ROOT, "out")
PNG = os.path.join(ROOT, "assets", "png")

os.makedirs(OUT, exist_ok=True)

MODULES = [
    "front", "part1", "part2", "part3", "part4", "part5", "part6", "revision",
]


def collect() -> str:
    chunks = []
    for name in MODULES:
        mod = __import__("content." + name, fromlist=["HTML"])
        chunks.append(mod.HTML)
    return "\n".join(chunks)


def write_html(body: str) -> str:
    css = os.path.join(ROOT, "book.css")
    html = f"""<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="utf-8">
<title>Java Concepts — Integrated Visual Notes</title>
<meta name="author" content="Java Concepts Integrated Visual Notes">
<link rel="stylesheet" href="{css}">
</head>
<body>
{body}
</body>
</html>
"""
    path = os.path.join(OUT, "Java_Concepts_Integrated_Visual_Notes.html")
    with open(path, "w", encoding="utf-8") as fh:
        fh.write(html)
    return path


def build_pdf(html_path: str) -> str:
    from weasyprint import HTML
    pdf = os.path.join(OUT, "Java_Concepts_Integrated_Visual_Notes.pdf")
    HTML(filename=html_path, base_url=ROOT).write_pdf(pdf)
    return pdf


# --------------------------------------------------------------------- DOCX
MARKERS = re.compile(r"[\x01-\x05]")
BACKTICK = re.compile(r"`+")
SECTION_NUM = re.compile(r"^(\d+(?:\.\d+)*)(?=[A-Z])")


def _plain(text: str) -> str:
    """Strip the converter's internal inline markers and inline-code backticks."""
    return BACKTICK.sub("", MARKERS.sub("", text))


class DocxConverter(HTMLParser):
    """Minimal, predictable HTML -> python-docx converter for this book."""
    def __init__(self, document, base_dir):
        super().__init__(convert_charrefs=True)
        self.d = document
        self.base = base_dir
        self.stack: list[str] = []
        self.buf: list[str] = []
        self.inline_bold = 0
        self.inline_italic = 0
        self.inline_code = 0
        self.pre_buf: list[str] = []
        self.in_pre = False
        self.list_depth = 0
        self.table_rows: list[list[str]] = []
        self.table_row: list[str] = []
        self.in_table = False
        self.in_cell = False
        self.cell_buf: list[str] = []
        self.skip = 0
        self.pending_image = None

    # -- helpers
    def flush_para(self, style=None):
        text = "".join(self.buf)
        text = re.sub(r"\s+", " ", text).strip()
        self.buf = []
        if not text:
            return
        from docx.shared import Pt
        p = self.d.add_paragraph()
        if style:
            p.style = self.d.styles[style]
        # honour **bold** markers captured as inline runs
        parts = re.split(r"(\x01|\x02|\x03)", self._raw)
        self._raw = ""
        self._emit_runs(p, text)

    _raw = ""

    def _emit_runs(self, p, text):
        from docx.shared import Pt
        # bold / italic / code markers encoded by handle_data via control chars
        tokens = re.split(r"(\x01|\x02|\x03|\x04|\x05)", text)
        bold = italic = code = False
        for tok in tokens:
            if tok == "\x01":
                bold = True; continue
            if tok == "\x02":
                bold = False; continue
            if tok == "\x03":
                italic = True; continue
            if tok == "\x04":
                italic = False; continue
            if tok == "\x05":
                code = not code; continue
            if not tok:
                continue
            r = p.add_run(tok)
            r.bold = bold
            r.italic = italic
            if code:
                r.font.name = "Consolas"
                r.font.size = Pt(9)
        return p

    # -- parser hooks
    def handle_starttag(self, tag, attrs):
        a = dict(attrs)
        if tag in ("h1", "h2", "h3", "h4", "p", "li", "td", "th", "figcaption",
                   "div", "span", "pre", "tr", "table", "ul", "ol", "body",
                   "head", "html", "figure", "img", "caption", "em", "strong",
                   "code", "b", "i"):
            self.stack.append(tag)
        if tag in ("h1", "h2", "h3", "h4"):
            self.buf = []
        elif tag == "pre":
            self.in_pre = True
            self.pre_buf = []
        elif tag == "strong" or tag == "b":
            self.buf.append("\x01")
        elif tag == "em" or tag == "i":
            self.buf.append("\x03")
        elif tag == "code":
            self.inline_code += 1
            self.buf.append("\x05")
        elif tag in ("ul", "ol"):
            self.list_depth += 1
        elif tag == "table":
            self.in_table = True
            self.table_rows = []
        elif tag == "tr":
            self.table_row = []
        elif tag in ("td", "th"):
            self.in_cell = True
            self.cell_buf = []
        elif tag == "img":
            self.pending_image = a.get("src", "")
        elif tag == "li":
            self.buf = []

    def handle_endtag(self, tag):
        if tag in ("strong", "b"):
            self.buf.append("\x02")
        elif tag in ("em", "i"):
            self.buf.append("\x04")
        elif tag == "code":
            self.inline_code = max(0, self.inline_code - 1)
            self.buf.append("\x05")
        elif tag == "pre":
            self.in_pre = False
            self._add_code("\n".join(self.pre_buf))
            self.pre_buf = []
        elif tag in ("h1", "h2", "h3", "h4"):
            level = int(tag[1])
            text = re.sub(r"\s+", " ", "".join(self.buf)).strip()
            self.buf = []
            if not text:
                return
            if self.pending_image:
                self._add_image(self.pending_image)
                self.pending_image = None
            if level == 1:
                self.d.add_page_break()
            self.d.add_heading(_plain(SECTION_NUM.sub(r"\1 ", text)), level=min(level, 4))
        elif tag == "p":
            if self.in_cell:
                return
            self._add_para()
        elif tag == "figcaption":
            self._add_para(italic=True, size=8)
        elif tag in ("td", "th"):
            self.in_cell = False
            self.table_row.append(_plain(re.sub(r"\s+", " ", "".join(self.cell_buf)).strip()))
            self.cell_buf = []
        elif tag == "tr":
            if self.table_row:
                self.table_rows.append(self.table_row)
            self.table_row = []
        elif tag == "table":
            self.in_table = False
            self._add_table()
        elif tag == "li":
            text = _plain(re.sub(r"\s+", " ", "".join(self.buf)).strip())
            self.buf = []
            if text:
                try:
                    self.d.add_paragraph(text, style="List Bullet")
                except KeyError:
                    self.d.add_paragraph("\u2022 " + text)
        elif tag == "img":
            if self.pending_image:
                self._add_image(self.pending_image)
                self.pending_image = None
        if self.stack and self.stack[-1] == tag:
            self.stack.pop()

    def handle_data(self, data):
        if self.in_pre:
            self.pre_buf.append(data.rstrip("\n"))
            return
        if self.in_cell:
            self.cell_buf.append(data)
            return
        self.buf.append(data)

    # -- emitters
    def _add_para(self, italic=False, size=None):
        from docx.shared import Pt
        raw = "".join(self.buf)
        self.buf = []
        # collapse whitespace per line, then join, so markers survive the collapse
        text = " ".join(part.strip() for part in raw.split("\n")).strip()
        text = re.sub(r"[ \t]{2,}", " ", text)
        if not _plain(text):
            return
        p = self.d.add_paragraph()
        self._emit_runs(p, text)
        if italic:
            for r in p.runs:
                r.italic = True
        if size:
            for r in p.runs:
                r.font.size = Pt(size)

    def _add_code(self, text):
        from docx.shared import Pt, RGBColor
        if not text.strip():
            return
        for line in text.split("\n"):
            p = self.d.add_paragraph()
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.space_before = Pt(0)
            r = p.add_run(line if line else " ")
            r.font.name = "Consolas"
            r.font.size = Pt(7.6)
            r.font.color.rgb = RGBColor(0x1A, 0x2A, 0x3D)

    def _add_table(self):
        rows = [r for r in self.table_rows if any(c for c in r)]
        if not rows:
            return
        ncols = max(len(r) for r in rows)
        t = self.d.add_table(rows=0, cols=ncols)
        t.style = "Light Grid Accent 1"
        for r in rows:
            cells = t.add_row().cells
            for i in range(ncols):
                cells[i].text = r[i] if i < len(r) else ""

    def _add_image(self, src):
        from docx.shared import Inches
        if not src:
            return
        path = src if os.path.isabs(src) else os.path.join(self.base, src)
        if not os.path.exists(path):
            return
        try:
            self.d.add_picture(path, width=Inches(6.4))
        except Exception:
            pass


def build_docx(html_path: str) -> str:
    from docx import Document
    from docx.shared import Pt, Inches
    from docx.enum.section import WD_SECTION

    with open(html_path, encoding="utf-8") as fh:
        html = fh.read()
    body = html.split("<body>", 1)[1].rsplit("</body>", 1)[0]

    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10.5)
    for s in doc.sections:
        s.left_margin = Inches(0.85)
        s.right_margin = Inches(0.85)
        s.top_margin = Inches(0.8)
        s.bottom_margin = Inches(0.8)

    conv = DocxConverter(doc, ROOT)
    conv.feed(body)

    path = os.path.join(OUT, "Java_Concepts_Integrated_Visual_Notes.docx")
    doc.save(path)
    return path


def main():
    what = sys.argv[1] if len(sys.argv) > 1 else "all"
    body = collect()
    html_path = write_html(body)
    print("HTML :", html_path, f"({os.path.getsize(html_path)//1024} KB)")
    if what in ("all", "pdf"):
        pdf = build_pdf(html_path)
        print("PDF  :", pdf, f"({os.path.getsize(pdf)//1024} KB)")
    if what in ("all", "docx"):
        docx = build_docx(html_path)
        print("DOCX :", docx, f"({os.path.getsize(docx)//1024} KB)")


if __name__ == "__main__":
    main()
